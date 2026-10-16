"""Quarterly real GDP by activity, from BCRD's reference-2018 workbook."""
from pathlib import Path
from datetime import datetime, timezone
from io import BytesIO
from urllib.parse import urljoin, urlparse
import json, math, re, unicodedata, requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data'/'gdp.json'
SOURCE='https://www.bancentral.gov.do/a/CustomView/2533-sector-real'
UA={'User-Agent':'Mozilla/5.0 (compatible; CFO-Macro-Dashboard/1.0)'}
PRINCIPAL=['Producto Interno Bruto','Agropecuario','Explotación de Minas y Canteras','Manufactura Local','Manufactura Zonas Francas','Construcción','Energía y Agua','Comercio','Hoteles, Bares y Restaurantes','Transporte y Almacenamiento','Comunicaciones','Intermediación Financiera, Seguros y Actividades Conexas','Actividades Inmobiliarias y de Alquiler','Enseñanza','Salud','Administración Pública y Defensa; Seguridad Social de Afiliación Obligatoria y Otros Servicios','Servicios profesionales','Otras Actividades de Servicios de Mercado']
AGGREGATES=['Industrias','Servicios','Valor Agregado','Impuestos a la producción netos de subsidios']
QUARTERS={'E-M':1,'A-J':2,'J-S':3,'O-D':4}

def clean(value):
    return re.sub(r'\s+',' ',''.join(c for c in unicodedata.normalize('NFKD',str(value)) if not unicodedata.combining(c))).strip().upper()
def key(label):return re.sub(r'[^a-z0-9]+','-',clean(label).lower()).strip('-')
def numeric(v):return isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v)

def periods(ws):
    rows=list(ws.iter_rows(values_only=True))
    header=next(i for i,r in enumerate(rows) if r[0] and clean(r[0]).startswith('INDICE DE VOLUMEN POR ACTIVIDAD'))
    year=None;out={}
    for i,v in enumerate(rows[header][1:],1):
        m=re.search(r'\b(20\d{2})\b',str(v))
        if m:year=m[1]
        q=QUARTERS.get(str(rows[header+1][i]).strip())
        if year and q:out[i]=f'{year}-Q{q}'
    if not out:raise ValueError('No se encontraron encabezados trimestrales')
    return out

def block(ws,prefix,columns):
    rows=list(ws.iter_rows(values_only=True));start=next(i for i,r in enumerate(rows) if r[0] and clean(r[0]).startswith(prefix))
    out={}
    for row in rows[start+2:]:
        if not row[0]:continue
        label=re.sub(r'\s+',' ',str(row[0])).strip();label_clean=clean(label)
        if any(label_clean.startswith(p) for p in ['INDICE DE VOLUMEN POR ACTIVIDAD','TASAS DE CRECIMIENTO POR ACTIVIDAD','INCIDENCIA POR ACTIVIDAD']):break
        values={p:float(row[i]) for i,p in columns.items() if i<len(row) and numeric(row[i])}
        if values:
            if key(label) in out:raise ValueError('Actividad duplicada en el bloque: '+label)
            out[key(label)]={'name':label,'values':values}
    return out

def parse_workbook(wb):
    ws=wb['PIBK_Trim'];cols=periods(ws)
    levels=block(ws,'INDICE DE VOLUMEN POR ACTIVIDAD',cols)
    rates=block(ws,'TASAS DE CRECIMIENTO POR ACTIVIDAD',cols)
    accum=block(wb['PIBK_Trim_Acum'],'TASAS DE CRECIMIENTO POR ACTIVIDAD',cols)
    principal=[key(x) for x in PRINCIPAL];aggregate={key(x) for x in AGGREGATES}
    missing=set(principal)-set(levels)
    if missing:raise ValueError('Faltan actividades principales: '+', '.join(sorted(missing)))
    published=sorted(levels[key('Producto Interno Bruto')]['values'])
    if len(published)<8:raise ValueError('Se necesitan ocho trimestres para comparar dos ventanas anuales')
    for a,b in zip(published,published[1:]):
        ordinal=lambda p:int(p[:4])*4+int(p[-1])-1
        if ordinal(b)!=ordinal(a)+1:raise ValueError('Trimestres incompletos en la serie')
    last=published[-4:];latest=last[-1];year=latest[:4];q=int(latest[-1]);items=[]
    for id,x in levels.items():
        if id in aggregate:continue
        v=x['values'];growth={}
        for p in published[4:]:
            prior=str(int(p[:4])-1)+p[4:]
            if p not in v or prior not in v or v[prior]<=0:raise ValueError('Índice faltante o inválido: '+x['name']+' '+p)
            calculated=(v[p]/v[prior]-1)*100
            reported=rates.get(id,{}).get('values',{}).get(p)
            if reported is None or abs(calculated-reported)>1e-6:raise ValueError('No concuerda la tasa oficial: '+x['name']+' '+p)
            growth[p]=reported
        previous=published[-8:-4]
        annual=(sum(v[p] for p in last)/sum(v[p] for p in previous)-1)*100
        ytd=accum.get(id,{}).get('values',{}).get(latest)
        current_ytd=[f'{year}-Q{i}' for i in range(1,q+1)]
        previous_ytd=[f'{int(year)-1}-Q{i}' for i in range(1,q+1)]
        expected=(sum(v[p] for p in current_ytd)/sum(v[p] for p in previous_ytd)-1)*100
        if ytd is None or abs(ytd-expected)>1e-6:raise ValueError('No concuerda el acumulado: '+x['name'])
        items.append({'id':id,'name':x['name'],'principal':id in principal,'yoy':growth,'index':v,'rolling_12m_yoy':annual,'ytd_yoy':ytd})
    order={id:i for i,id in enumerate(principal)}
    items.sort(key=lambda x:(0,order[x['id']]) if x['principal'] else (1,x['name']))
    return {'latest_period':latest,'periods':last,'series':items,'history_start':published[0],'quarter_count':len(published),'reference_year':2018,'preliminary':True,'ytd_period':{'year':int(year),'through_quarter':q},'rolling_periods':last,'rolling_comparison_periods':previous}

def discover():
    r=requests.get(SOURCE,headers=UA,timeout=35);r.raise_for_status();soup=BeautifulSoup(r.text,'html.parser')
    for a in soup.find_all('a',href=True):
        row=a.find_parent('tr');context=clean(row.get_text(' ',strip=True)) if row else ''
        url=urljoin(SOURCE,a['href'])
        if ('PIB_ORIGEN_2018.XLSX' in url.upper() or ('POR SECTORES DE ORIGEN' in context and '2018' in context and 'HOMOGENEA' not in context)) and urlparse(url).path.lower().endswith('.xlsx'):
            host=urlparse(url).hostname or ''
            if host=='bancentral.gov.do' or host.endswith('.bancentral.gov.do'):return url
    raise ValueError('No se encontró el Excel vigente de PIB por actividad, referencia 2018')

def main():
    now=datetime.now(timezone.utc).isoformat(timespec='seconds')
    try:
        url=discover();r=requests.get(url,headers=UA,timeout=60);r.raise_for_status()
        wb=load_workbook(BytesIO(r.content),read_only=True,data_only=True)
        try:payload=parse_workbook(wb)
        finally:wb.close()
        if DATA.exists():
            previous=json.loads(DATA.read_text())
            if payload['latest_period']<previous.get('latest_period',''):raise ValueError('La fuente devolvió un trimestre anterior al vigente')
        payload.update(source='BCRD',source_page=SOURCE,source_url=url,verified_at=now,checked_at=now,source_health={'ok':True,'error':None})
        DATA.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n')
        print('GDP:',payload['latest_period'],len(payload['series']),'activities;',payload['quarter_count'],'quarters')
    except Exception as exc:
        payload=json.loads(DATA.read_text()) if DATA.exists() else {'series':[],'periods':[]}
        payload.update(checked_at=now,source_health={'ok':False,'error':str(exc)})
        DATA.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n')
        print('GDP verification incomplete:',exc)

if __name__=='__main__':main()
