"""Public input-market signals only. No purchasing, inventory or supplier records."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone, date
from io import StringIO, BytesIO
from pathlib import Path
import csv
import json
import math
import re
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from input_catalog import catalogue

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data' / 'inputs.json'
HEADERS = {'User-Agent': 'Mozilla/5.0 (compatible; Public-Input-Market-Monitor/1.0)'}
GSCPI_PAGE = 'https://www.newyorkfed.org/research/policy/gscpi'
FDA_URL = 'https://api.fda.gov/drug/shortages.json'

# Macro references are not quotes for pharmaceutical grades or converted packaging.
SERIES = {
    'aluminum': dict(name='Aluminio primario', series_id='WB_ALUMINUM', source='Banco Mundial · Pink Sheet', unit='USD / tonelada', scope='Aluminio primario LME, referencia del Banco Mundial; no es precio de foil farmacéutico.', category='empaque'),
    'resins': dict(name='Resinas termoplásticas', series_id='WPU0662', source='BLS · fuente directa', unit='Índice PPI', scope='Canasta de resinas en EE. UU.; señal amplia para plásticos. No mide específicamente PVC, PVDC ni PET farmacéutico.', category='empaque'),
    'corrugated': dict(name='Cartón corrugado', series_id='PCU32221132221105', source='BLS · fuente directa', unit='Índice PPI', scope='Láminas y rollos de cartón corrugado en EE. UU.; no es una cotización de cajas en RD.', category='empaque'),
    'paperboard': dict(name='Cajas plegadizas', series_id='WPU09150336', source='BLS · fuente directa', unit='Índice PPI', scope='Cajas, empaques y componentes de cartón plegadizo en EE. UU.; referencia externa.', category='empaque'),
    'ethanol': dict(name='Etanol', series_id='WPU06140341', source='BLS · fuente directa', unit='Índice PPI', scope='Etanol en EE. UU.; no distingue concentración, grado farmacéutico ni precio local. No representa isopropanol.', category='excipiente'),
    'sugar': dict(name='Azúcar mundial', series_id='WB_SUGAR_WORLD', source='Banco Mundial · Pink Sheet', unit='Centavos USD / libra', scope='Precio mundial del azúcar; referencia indirecta para sacarosa. No cotiza sacarosa farmacéutica, mezclas de compresión directa ni polioles.', category='excipiente'),
}


def parse_fred_csv(content, series_id, today=None):
    today = today or date.today()
    rows = csv.DictReader(StringIO(content))
    if rows.fieldnames not in (['observation_date', series_id], ['DATE', series_id]):
        raise ValueError('Encabezado de serie no válido')
    dates = {}
    for row in rows:
        raw_date = row.get('observation_date', row.get('DATE', ''))
        if not re.fullmatch(r'\d{4}-\d{2}-01', raw_date or ''):
            raise ValueError('Se esperaba una frecuencia mensual')
        day = date.fromisoformat(raw_date)
        if day > today or row[series_id] in ('', '.', None):
            continue
        value = float(row[series_id])
        if not math.isfinite(value) or value <= 0:
            raise ValueError('Observación fuera de rango')
        dates[raw_date[:7]] = value
    if not dates:
        raise ValueError('Sin observaciones mensuales válidas')
    return [{'period': p, 'value': round(v, 6)} for p, v in sorted(dates.items())][-61:]


def month_before(period, months):
    y, m = map(int, period.split('-'))
    serial = y * 12 + m - 1 - months
    return f'{serial // 12:04d}-{serial % 12 + 1:02d}'


def changes(history):
    if not history:
        return {key: None for key in ('mom', 'qoq', 'yoy')}
    values = {x['period']: x['value'] for x in history}
    latest = history[-1]
    result = {}
    for key, months in [('mom', 1), ('qoq', 3), ('yoy', 12)]:
        base = values.get(month_before(latest['period'], months))
        result[key] = round((latest['value'] / base - 1) * 100, 2) if base is not None and base > 0 else None
    return result


def get(url, **kwargs):
    r = requests.get(url, headers=HEADERS, timeout=25, **kwargs)
    r.raise_for_status()
    return r


BLS_API = 'https://api.bls.gov/publicAPI/v1/timeseries/data/'
WB_PAGE = 'https://www.worldbank.org/en/research/commodity-markets'
BLS_KEYS = ('resins', 'corrugated', 'paperboard', 'ethanol')
WB_KEYS = ('aluminum', 'sugar')

def signal(key, history, url, **metadata):
    return dict(SERIES[key], id=key, url=url, history=history, **history[-1], changes=changes(history), **metadata)

def parse_bls(payload, series_id, today=None):
    today = today or date.today()
    if payload.get('status') != 'REQUEST_SUCCEEDED':
        raise ValueError('BLS no completó la consulta')
    matches = [x for x in payload.get('Results', {}).get('series', []) if x.get('seriesID') == series_id]
    if len(matches) != 1:raise ValueError('BLS no devolvió la serie solicitada')
    values = {}
    for x in matches[0].get('data', []):
        if not re.fullmatch(r'M(?:0[1-9]|1[0-2])', x.get('period', '')):continue
        period = f"{int(x['year']):04d}-{x['period'][1:]}"
        if period > today.strftime('%Y-%m') or str(x.get('value', '')).strip() in ('', '-', '.', 'NA'):continue
        value = float(x['value'])
        if not math.isfinite(value) or value <= 0:raise ValueError('Índice BLS fuera de rango')
        values[period] = value
    if not values:raise ValueError('Sin meses BLS válidos')
    return [{'period':p,'value':v} for p,v in sorted(values.items())][-61:]

def bls():
    # One unregistered request for all four series stays well below BLS limits.
    year=date.today().year
    r=requests.post(BLS_API,headers=HEADERS,json={'seriesid':[SERIES[k]['series_id'] for k in BLS_KEYS],'startyear':str(year-5),'endyear':str(year)},timeout=30)
    r.raise_for_status();payload=r.json();results={}
    for key in BLS_KEYS:
        try:
            history=parse_bls(payload,SERIES[key]['series_id'])
            results[key]=(signal(key,history,'https://data.bls.gov/timeseries/'+SERIES[key]['series_id']),None)
        except (ValueError,KeyError,TypeError) as exc:results[key]=(None,str(exc))
    return results

def parse_worldbank(wb, today=None):
    today=today or date.today()
    if 'Monthly Prices' not in wb.sheetnames:raise ValueError('Falta la hoja de precios mensuales')
    rows=iter(wb['Monthly Prices'].values);headers=None
    for row in rows:
        if 'Aluminum' in row and 'Sugar, world' in row:
            headers=row;break
    if headers is None:raise ValueError('Faltan columnas de aluminio y azúcar mundial')
    cols={'aluminum':headers.index('Aluminum'),'sugar':headers.index('Sugar, world')}
    units=next(rows)
    if units[cols['aluminum']]!='($/mt)' or units[cols['sugar']]!='($/kg)':raise ValueError('Unidades del Banco Mundial no reconocidas')
    histories={k:{} for k in WB_KEYS}
    for row in rows:
        m=re.fullmatch(r'(\d{4})M(0[1-9]|1[0-2])',str(row[0]))
        if not m:continue
        period=f'{m[1]}-{m[2]}'
        if period>today.strftime('%Y-%m'):continue
        for key,col in cols.items():
            raw=row[col]
            if raw is None or str(raw).strip() in ('…','...','','NA'):continue
            value=float(raw)
            if not math.isfinite(value) or value<=0:raise ValueError('Precio del Banco Mundial fuera de rango')
            # World Bank sugar is US$/kg. Keep displayed US cents/lb consistently.
            if key=='sugar':value*=100*0.45359237
            histories[key][period]=value
    if any(not h for h in histories.values()):raise ValueError('Sin histórico comparable del Banco Mundial')
    return {k:[{'period':p,'value':round(v,6)} for p,v in sorted(h.items())][-61:] for k,h in histories.items()}

def worldbank():
    # Discover the current workbook; archived year-specific URLs can stay frozen.
    soup=BeautifulSoup(get(WB_PAGE).text,'html.parser')
    links=[a['href'] for a in soup.find_all('a',href=True) if a['href'].endswith('/CMO-Historical-Data-Monthly.xlsx') and a['href'].startswith('https://thedocs.worldbank.org/')]
    if not links:raise ValueError('No se encontró el archivo mensual vigente del Banco Mundial')
    url=links[0]
    wb=load_workbook(BytesIO(get(url).content),read_only=True,data_only=True)
    try:histories=parse_worldbank(wb)
    finally:wb.close()
    return {k:(signal(k,h,WB_PAGE,download_url=url,methodology='Histórico completo reconstruido con Banco Mundial; no se empalma con FMI/FRED.'),None) for k,h in histories.items()}


def parse_gscpi(content, today=None):
    today = today or date.today()
    rows = csv.DictReader(StringIO(content))
    if not rows.fieldnames or rows.fieldnames[0] != 'Date':
        raise ValueError('Encabezado GSCPI no identificado')
    vintages = []
    for name in rows.fieldnames[1:]:
        try:
            vintage = datetime.strptime(name, '%b-%y').date()
            if vintage <= today:
                vintages.append((vintage, name))
        except ValueError:
            continue
    if not vintages:
        raise ValueError('Sin edición GSCPI identificada')
    column = max(vintages)[1]
    observations = {}
    for row in rows:
        if not row.get('Date', '').strip():
            continue
        day = datetime.strptime(row['Date'], '%d-%b-%Y').date()
        raw = row.get(column, '')
        if day > today or not raw or raw.strip().lower() in ('na', 'n/a', '.'):
            continue
        value = float(raw)
        if not math.isfinite(value) or abs(value) > 20:
            raise ValueError('GSCPI fuera de rango')
        observations[day.strftime('%Y-%m')] = value
    if not observations:
        raise ValueError('Sin observaciones GSCPI válidas')
    return [{'period': p, 'value': v} for p, v in sorted(observations.items())][-61:]


def gscpi():
    url = 'https://www.newyorkfed.org/medialibrary/research/interactives/data/gscpi/gscpi_interactive_data.csv'
    history = parse_gscpi(get(url).text)
    return dict(id='gscpi', name='Presión logística global', source='Federal Reserve Bank of New York', url=GSCPI_PAGE, unit='Desviaciones estándar', category='logistica', scope='GSCPI: presión frente al promedio histórico. No mide disponibilidad de una sustancia ni una ruta a RD. Sus variaciones se leen en puntos, no en porcentajes.', history=history, **history[-1])


def matches_active(row, active):
    text = row.get('generic_name', '').casefold()
    def has(term):
        return re.search(r'(?<![a-z])' + re.escape(term) + r'(?![a-z])', text) is not None
    return all(has(t) for t in active['terms']) and not any(has(t) for t in active['exclude'])


def parse_shortages(rows, catalog):
    items = []
    for active in catalog['actives']:
        current = [r for r in rows if str(r.get('status', '')).casefold() == 'current' and matches_active(r, active)]
        # FDA records describe finished-drug presentations, not API stock.
        products = {}
        for row in current:
            record = {k: row.get(k, '') for k in ('generic_name', 'presentation', 'company_name', 'availability', 'shortage_reason', 'update_date')}
            key = (record['presentation'], record['company_name'])
            products[key] = record
        items.append({'name': active['name'], 'count': len(products), 'presentations': list(products.values())})
    return items


def shortages(catalog):
    rows, total, updated = [], None, None
    for skip in range(0, 5000, 1000):
        payload = get(FDA_URL, params={'limit': 1000, 'skip': skip}).json()
        meta = payload.get('meta', {})
        count = meta.get('results', {}).get('total')
        batch = payload.get('results')
        if not isinstance(count, int) or not isinstance(batch, list) or not batch:
            raise ValueError('Respuesta FDA incompleta')
        if total is None:
            total, updated = count, meta.get('last_updated')
        elif count != total or meta.get('last_updated') != updated:
            raise ValueError('La base FDA cambió durante la consulta')
        rows.extend(batch)
        if len(rows) >= total:
            break
    if total is None or len(rows) != total:
        raise ValueError('No se recuperó toda la base FDA')
    if not updated or date.fromisoformat(updated) > date.today():
        raise ValueError('Fecha de base FDA no válida')
    return {'last_updated': updated, 'url': 'https://www.fda.gov/drugs/drug-safety-and-availability/drug-shortages', 'source': 'FDA · openFDA Drug Shortages', 'items': parse_shortages(rows, catalog)}


def merge_series(previous, results, stamp):
    series, health = {}, {}
    for key, (observation, error) in results.items():
        health[key] = {'ok': error is None, 'checked_at': stamp, 'error': error}
        if observation is not None:
            series[key] = dict(observation, verified_at=stamp)
        elif key in previous:
            # A failed check must not refresh the last successful verification.
            series[key] = previous[key]
    return series, health


def main():
    stamp = datetime.now(timezone.utc).isoformat(timespec='seconds')
    try:
        previous = json.loads(DATA.read_text())
    except (FileNotFoundError, ValueError):
        previous = {}

    catalog = catalogue()
    def run(group):
        try:
            if group=='bls':return bls()
            if group=='worldbank':return worldbank()
            value=shortages(catalog) if group=='fda' else gscpi()
            return {group:(value,None)}
        except Exception as exc:
            keys=BLS_KEYS if group=='bls' else WB_KEYS if group=='worldbank' else (group,)
            return {k:(None,type(exc).__name__+': '+str(exc)[:250]) for k in keys}

    with ThreadPoolExecutor(max_workers=4) as pool:
        results={k:v for group in pool.map(run,['bls','worldbank','gscpi','fda']) for k,v in group.items()}
    for key,(observation,error) in list(results.items()):
        previous_period=previous.get('series',{}).get(key,{}).get('period','')
        if key!='fda' and observation and observation.get('period','')<previous_period:
            results[key]=(None,'La fuente devolvió un período anterior; se conserva el último histórico verificado.')
    fda_result, fda_error = results.pop('fda')
    series, health = merge_series(previous.get('series', {}), results, stamp)
    health['fda'] = {'ok': fda_error is None, 'checked_at': stamp, 'error': fda_error}
    fda_data = dict(fda_result, verified_at=stamp) if fda_result else previous.get('fda')
    payload = {'checked_at': stamp, 'series': series, 'source_health': health, 'catalogue': catalog, 'fda': fda_data}
    DATA.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n')
    print('Public input signals:', {k: v['ok'] for k, v in health.items()})


if __name__ == '__main__':
    main()
