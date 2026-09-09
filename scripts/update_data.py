"""CFO Macro RD — public data updater v4.
Updates validated BCRD, DGII and Hacienda indicators and generates CFO Watch rules.
Keeps the last validated value whenever a field cannot be parsed.
"""
from pathlib import Path
from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZipFile
import json, re, requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "current.json"
BCRD = "https://www.bancentral.gov.do/"
HACIENDA = "https://www.hacienda.gob.do/"
DIGEPRES_EXECUTION = "https://digepres.gob.do/ejecucion-presupuestaria-{year}/"
DGII_PAGE = "https://www.dgii.gov.do/estadisticas/Operaciones-Recaudaciones-ITBIS/Paginas/default.aspx"
DGII_FILES = {
    "itbis_total_operations": "https://www.dgii.gov.do/estadisticas/Operaciones-Recaudaciones-ITBIS/Operaciones%20Totales/Totales-ITBIS-2015-2026.zip",
    "itbis_taxed_operations": "https://www.dgii.gov.do/estadisticas/Operaciones-Recaudaciones-ITBIS/Operaciones%20Gravadas/Gravadas-ITBIS-2015-2026.zip",
    "itbis_revenue": "https://www.dgii.gov.do/estadisticas/Operaciones-Recaudaciones-ITBIS/Recaudacion%20ITBIS/Recaudacion-ITBIS-2015-2026.zip",
}
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36"}
MONTHS = {"enero":"01","febrero":"02","marzo":"03","abril":"04","mayo":"05","junio":"06",
          "julio":"07","agosto":"08","septiembre":"09","octubre":"10","noviembre":"11","diciembre":"12"}

def norm(s): return re.sub(r"\s+", " ", s.replace("\xa0"," ")).strip()
def per(month, year): return f"{year}-{MONTHS[month.lower()]}"
def num(s): return float(s.replace(",", ""))

MONTH_NAMES = {v: k[:3] for k, v in MONTHS.items()}

def dgii_workbook(url):
    headers = dict(UA)
    headers["Referer"] = DGII_PAGE
    r = requests.get(url, headers=headers, timeout=60)
    r.raise_for_status()
    with ZipFile(BytesIO(r.content)) as z:
        names = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
        if not names:
            raise ValueError("El ZIP de DGII no contiene un archivo XLSX")
        return load_workbook(BytesIO(z.read(names[0])), read_only=True, data_only=True)

def dgii_monthly_total(ws, year):
    """Return monthly values from the row named Total, regardless of header row shifts."""
    header_row = None
    months = {}
    for row_number, row in enumerate(ws.iter_rows(), 1):
        vals = [c.value for c in row]
        for i, value in enumerate(vals):
            if isinstance(value, str) and value.strip().lower() in MONTHS:
                months[i] = MONTHS[value.strip().lower()]
        if months:
            header_row = row_number
            break
    if header_row is None:
        raise ValueError(f"No se encontraron meses en la hoja {year}")
    for row in ws.iter_rows(min_row=header_row + 1):
        vals = [c.value for c in row]
        label = next((v for v in vals[:3] if isinstance(v, str) and v.strip()), "")
        if label.strip().lower() == "total":
            out = {}
            for col, month in months.items():
                value = vals[col] if col < len(vals) else None
                if isinstance(value, (int, float)):
                    out[f"{year}-{month}"] = float(value) / 1_000_000
            if not out:
                raise ValueError(f"La fila Total de {year} no contiene valores mensuales")
            return out
    raise ValueError(f"No se encontró la fila Total en la hoja {year}")

def update_dgii(p, hist):
    updated, errors = [], {}
    dgii_hist = hist.setdefault("monthly", {})
    for key, url in DGII_FILES.items():
        try:
            wb = dgii_workbook(url)
            values = {}
            for year in ("2025", "2026"):
                if year in wb.sheetnames:
                    values.update(dgii_monthly_total(wb[year], year))
            if not values:
                raise ValueError("No se extrajeron observaciones")
            dgii_hist.setdefault(key, {}).update(values)
            latest_period = max(values)
            latest_value = values[latest_period]
            prior_period = f"{int(latest_period[:4])-1}{latest_period[4:]}"
            prior_value = dgii_hist[key].get(prior_period)
            yoy = ((latest_value / prior_value) - 1) * 100 if prior_value else None
            setv(p, key, latest_value, latest_period, "DGII")
            if yoy is not None:
                setv(p, key + "_yoy", yoy, latest_period, "DGII")
            updated.append(key)
        except Exception as exc:
            errors[key] = str(exc)

    total = p.get("series", {}).get("itbis_total_operations", {})
    taxed = p.get("series", {}).get("itbis_taxed_operations", {})
    if total.get("period") == taxed.get("period") and total.get("value"):
        setv(p, "itbis_taxed_share", taxed["value"] / total["value"] * 100, total["period"], "DGII")
    p.setdefault("source_health", {})["dgii"] = {
        "ok": not errors,
        "url": DGII_PAGE,
        "updated_fields": updated,
        "errors": errors,
    }
    return updated

def update_hacienda(p, hist):
    """Update the headline fiscal execution figures published by Hacienda."""
    updated = []
    try:
        r = requests.get(HACIENDA, headers=UA, timeout=30)
        r.raise_for_status()
        text = norm(BeautifulSoup(r.text, "html.parser").get_text(" ", strip=True))
        pattern = (
            r"Indicadores fiscales\s+Ingresos\s+([\d,]+(?:\.\d+)?)\s+"
            r"Gastos\s+([\d,]+(?:\.\d+)?)\s+Fuentes Financieras\s+"
            r"([\d,]+(?:\.\d+)?)\s+Aplicaciones Financieras\s+"
            r"([\d,]+(?:\.\d+)?).*?Actualizado al\s+(\d{1,2})\s+"
            r"(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre),\s+(\d{4})"
        )
        m = re.search(pattern, text, re.I | re.S)
        if not m:
            raise ValueError("No se encontró el bloque de indicadores fiscales")
        income, expense, sources, applications = map(num, m.groups()[:4])
        day, month, year = m.groups()[4:]
        period = f"{year}-{MONTHS[month.lower()]}-{int(day):02d}"
        values = {
            "fiscal_income": income,
            "fiscal_expense": expense,
            "financial_sources": sources,
            "financial_applications": applications,
            "budget_balance_simple": income - expense,
            "net_financing": sources - applications,
        }
        fiscal_hist = hist.setdefault("fiscal", {})
        for key, value in values.items():
            if setv(p, key, value, period, "Hacienda"):
                updated.append(key)
            fiscal_hist.setdefault(key, {})[period] = value
        p.setdefault("source_health", {})["hacienda"] = {
            "ok": True,
            "url": HACIENDA,
            "http_status": r.status_code,
            "updated_fields": updated,
        }
    except Exception as exc:
        p.setdefault("source_health", {})["hacienda"] = {
            "ok": False,
            "url": HACIENDA,
            "updated_fields": [],
            "error": str(exc),
        }
    return updated

def find_fiscal_summary(wb):
    """Extract comparable monthly fiscal results from a DIGEPRES workbook."""
    for ws in wb.worksheets:
        title_row = None
        title_text = ""
        for row_number, row in enumerate(ws.iter_rows(max_row=min(ws.max_row, 45)), 1):
            values = [c.value for c in row]
            joined = " ".join(str(v) for v in values if v is not None)
            if "Resultados Presupuestarios" in joined:
                title_row = row_number
                title_text = joined
                break
        if title_row is None:
            continue

        year_columns = {}
        for row in ws.iter_rows(min_row=title_row, max_row=min(title_row + 8, ws.max_row)):
            for index, cell in enumerate(row):
                value = cell.value
                if isinstance(value, int) and 2010 <= value <= 2035:
                    year_columns[index] = str(value)
                elif isinstance(value, str) and re.fullmatch(r"20\d{2}", value.strip()):
                    year_columns[index] = value.strip()
        if not year_columns:
            title_years = re.findall(r"20\d{2}", title_text)
            current_year = title_years[-1] if title_years else None
            if current_year:
                for row in ws.iter_rows(min_row=title_row, max_row=min(title_row + 8, ws.max_row)):
                    for index, cell in enumerate(row):
                        if isinstance(cell.value, str) and norm(cell.value).lower() == "devengado":
                            year_columns[index] = current_year
            if not year_columns:
                continue

        rows = {}
        aliases = {
            "1 - ingresos": "fiscal_income",
            "2 - gastos": "fiscal_expense",
            "resultado financiero": "budget_balance_simple",
            "financiamiento neto": "net_financing",
        }
        for row in ws.iter_rows(min_row=title_row + 1, max_row=min(title_row + 35, ws.max_row)):
            values = [c.value for c in row]
            label = next((norm(v).lower() for v in values if isinstance(v, str) and norm(v)), "")
            for prefix, key in aliases.items():
                if label.startswith(prefix):
                    rows[key] = values
                    break
        if "fiscal_income" not in rows or "fiscal_expense" not in rows:
            continue

        observations = {}
        for column, year in year_columns.items():
            item = {}
            for key, values in rows.items():
                value = values[column] if column < len(values) else None
                if isinstance(value, (int, float)):
                    item[key] = float(value) / 1_000_000 if abs(value) > 10_000_000 else float(value)
            if "fiscal_income" in item and "fiscal_expense" in item:
                item.setdefault("budget_balance_simple", item["fiscal_income"] - item["fiscal_expense"])
                observations[year] = item

        budget = None
        current_year = max(observations) if observations else None
        if current_year:
            income_values = rows["fiscal_income"]
            candidate_columns = [c for c in range(min(year_columns)) if c < len(income_values)]
            candidates = [income_values[c] for c in candidate_columns if isinstance(income_values[c], (int, float))]
            if candidates:
                raw = candidates[-1]
                budget = float(raw) / 1_000_000 if abs(raw) > 10_000_000 else float(raw)
        return observations, budget
    raise ValueError("No se encontró la tabla de resultados presupuestarios")

def fiscal_month_from_url(url):
    name = url.rsplit("/", 1)[-1].lower()
    if any(x in name for x in ("enero-junio", "enero-marzo", "enero-septiembre", "enero-diciembre", "avance", "completo")):
        return None
    found = [(number, month) for month, number in MONTHS.items() if month in name]
    years = re.findall(r"20\d{2}", name)
    if not found or not years:
        return None
    return years[-1], found[-1][0]

def update_fiscal_history(p, hist):
    """Build comparable monthly fiscal series from DIGEPRES Excel reports."""
    fiscal = hist.setdefault("fiscal_monthly", {})
    errors = {}
    parsed = 0
    # 2025 reports supply 2024 comparatives where available; 2026 reports supply
    # 2025 comparatives. Limit discovery to these pages so routine runs stay fast.
    for publication_year in (2025, 2026):
        try:
            page = requests.get(DIGEPRES_EXECUTION.format(year=publication_year), headers=UA, timeout=40)
            page.raise_for_status()
            soup = BeautifulSoup(page.text, "html.parser")
            links = []
            for a in soup.find_all("a", href=True):
                url = a["href"]
                target = fiscal_month_from_url(url)
                if target and url.lower().endswith((".xlsx", ".xls")) and url not in links:
                    links.append(url)
            for url in links:
                target_year, target_month = fiscal_month_from_url(url)
                period = f"{target_year}-{target_month}"
                if period in fiscal.get("fiscal_income", {}):
                    continue
                try:
                    response = requests.get(url, headers=UA, timeout=90)
                    response.raise_for_status()
                    wb = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
                    observations, budget = find_fiscal_summary(wb)
                    for year, values in observations.items():
                        obs_period = f"{year}-{target_month}"
                        for key, value in values.items():
                            fiscal.setdefault(key, {})[obs_period] = value
                    if budget is not None:
                        fiscal.setdefault("annual_income_budget", {})[target_year] = budget
                    parsed += 1
                except Exception as exc:
                    errors[url.rsplit("/", 1)[-1]] = str(exc)
        except Exception as exc:
            errors[str(publication_year)] = str(exc)

    periods = sorted(fiscal.get("fiscal_income", {}))
    if periods:
        latest = periods[-1]
        previous = f"{int(latest[:4])-1}{latest[4:]}"
        for key in ("fiscal_income", "fiscal_expense"):
            current_value = fiscal.get(key, {}).get(latest)
            prior_value = fiscal.get(key, {}).get(previous)
            if current_value is not None and prior_value:
                setv(p, key + "_yoy", (current_value / prior_value - 1) * 100, latest, "DIGEPRES")
        current_balance = fiscal.get("budget_balance_simple", {}).get(latest)
        prior_balance = fiscal.get("budget_balance_simple", {}).get(previous)
        if current_balance is not None and prior_balance is not None:
            setv(p, "budget_balance_yoy_change", current_balance - prior_balance, latest, "DIGEPRES")

        year, month = latest.split("-")
        ytd_income = sum(v for k, v in fiscal.get("fiscal_income", {}).items() if k.startswith(year + "-") and k <= latest)
        budget = fiscal.get("annual_income_budget", {}).get(year)
        if budget:
            setv(p, "fiscal_income_execution_pct", ytd_income / budget * 100, latest, "DIGEPRES")
            setv(p, "fiscal_expected_execution_pct", int(month) / 12 * 100, latest, "Cálculo")

    p.setdefault("source_health", {})["digepres_history"] = {
        "ok": bool(periods),
        "url": DIGEPRES_EXECUTION.format(year=2026),
        "parsed_files": parsed,
        "latest_period": periods[-1] if periods else None,
        "errors": errors,
    }

def setv(p,k,v,period,source="BCRD"):
    if v is None or period is None: return False
    p.setdefault("series",{})[k]={"value":v,"period":period,"source":source}
    return True

def fx(text):
    m=re.search(r"Tipo de cambio\s+(\d{1,2})\s+de\s+(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre)\s+(\d{4}).{0,140}?Compra\s+(\d{1,3}\.\d{4}).{0,100}?Venta\s+(\d{1,3}\.\d{4})",text,re.I|re.S)
    if not m:return None
    d,mo,y,b,s=m.groups(); b=float(b); s=float(s)
    if not (30<b<100 and 30<s<100 and s>=b): return None
    return b,s,f"{y}-{MONTHS[mo.lower()]}-{int(d):02d}"

def watch(p, hist):
    s=p["series"]; items=[]
    inf=s.get("inflation_yoy",{}).get("value")
    if inf is not None:
        if inf>5:
            items.append({"level":"yellow","title":"Inflación: vigilancia","text":f"La inflación interanual se ubica en {inf:.2f}%, por encima del techo de 5.0% del rango meta del BCRD."})
        else:
            items.append({"level":"green","title":"Inflación: dentro de meta","text":f"La inflación interanual se ubica en {inf:.2f}%, dentro del rango meta de 3%-5%."})
    a=s.get("active_rate",{}).get("value"); pol=s.get("policy_rate",{}).get("value")
    if a is not None and pol is not None:
        items.append({"level":"yellow" if a-pol>=8 else "green","title":"Financiamiento: spread bancario","text":f"La tasa activa promedio es {a:.2f}% frente a una TPM de {pol:.2f}%, un diferencial de {a-pol:.2f} puntos porcentuales."})
    im=s.get("imae_yoy",{}).get("value")
    if im is not None:
        items.append({"level":"green" if im>=4 else "yellow","title":"Actividad económica","text":f"El IMAE registra crecimiento interanual de {im:.1f}% en el último dato publicado."})
    r=s.get("gross_reserves",{}).get("value")
    if r is not None:
        items.append({"level":"green","title":"Sector externo: reservas","text":f"Las reservas internacionales brutas se sitúan en US${r:,.1f} millones."})

    # DGII: real activity, sustained deceleration and collection divergence.
    dgii_alert_count=len(items)
    ops_yoy=s.get("itbis_total_operations_yoy",{}).get("value")
    if ops_yoy is not None and inf is not None and ops_yoy < inf:
        items.append({"level":"red","title":"DGII: actividad real debilitándose","text":f"Las operaciones declaradas crecen {ops_yoy:.1f}% interanual, por debajo de la inflación de {inf:.1f}%; implica una contracción real aproximada."})
    monthly=hist.get("monthly",{})
    ops=monthly.get("itbis_total_operations",{})
    op_periods=sorted(ops)
    yoy_path=[]
    for period in op_periods[-5:]:
        prior=f"{int(period[:4])-1}{period[4:]}"
        if ops.get(prior): yoy_path.append((ops[period]/ops[prior]-1)*100)
    if len(yoy_path)>=3 and yoy_path[-3] > yoy_path[-2] > yoy_path[-1]:
        items.append({"level":"yellow","title":"DGII: desaceleración sostenida","text":f"El crecimiento interanual de las operaciones se ha moderado durante tres meses consecutivos ({yoy_path[-3]:.1f}% → {yoy_path[-2]:.1f}% → {yoy_path[-1]:.1f}%)."})
    revenue_yoy=s.get("itbis_revenue_yoy",{}).get("value")
    if revenue_yoy is not None and ops_yoy is not None and revenue_yoy-ops_yoy>=5:
        items.append({"level":"yellow","title":"DGII: divergencia tributaria","text":f"La recaudación ITBIS crece {revenue_yoy:.1f}% frente a {ops_yoy:.1f}% en operaciones; la brecha puede reflejar fiscalización, composición o pagos extraordinarios."})
    if len(items)==dgii_alert_count and ops_yoy is not None:
        real_gap=ops_yoy-inf if inf is not None else None
        detail=f", {real_gap:.1f} puntos por encima de la inflación" if real_gap is not None else ""
        items.append({"level":"green","title":"DGII: actividad formal estable","text":f"Las operaciones declaradas crecen {ops_yoy:.1f}% interanual{detail}. No se activan alertas de desaceleración real ni divergencia tributaria."})

    # Hacienda/DIGEPRES: comparable monthly fiscal signals.
    income_yoy=s.get("fiscal_income_yoy",{}).get("value")
    expense_yoy=s.get("fiscal_expense_yoy",{}).get("value")
    if income_yoy is not None and expense_yoy is not None and income_yoy < expense_yoy:
        items.append({"level":"yellow","title":"Hacienda: gasto supera ingresos","text":f"El gasto mensual crece {expense_yoy:.1f}% interanual frente a {income_yoy:.1f}% de los ingresos, una brecha de {expense_yoy-income_yoy:.1f} puntos."})
    balance_change=s.get("budget_balance_yoy_change",{}).get("value")
    if balance_change is not None and balance_change < 0:
        items.append({"level":"red","title":"Hacienda: deterioro fiscal","text":f"El balance mensual se deterioró en RD${abs(balance_change):,.0f} MM frente al mismo mes del año anterior."})
    execution=s.get("fiscal_income_execution_pct",{}).get("value")
    expected=s.get("fiscal_expected_execution_pct",{}).get("value")
    if execution is not None and expected is not None and execution < expected-5:
        items.append({"level":"yellow","title":"Hacienda: ingresos rezagados","text":f"La ejecución acumulada de ingresos es {execution:.1f}% del presupuesto, frente a una referencia temporal de {expected:.1f}%."})
    p["cfo_watch"]=items

def main():
    p=json.loads(DATA.read_text(encoding="utf-8"))
    r=requests.get(BCRD,headers=UA,timeout=30); r.raise_for_status()
    text=norm(BeautifulSoup(r.text,"html.parser").get_text(" ",strip=True))
    u=[]
    z=fx(text)
    if z:
        b,s,period=z
        for k,v in [("usd_dop_buy",b),("usd_dop_sell",s)]:
            if setv(p,k,v,period):u.append(k)

    patterns=[
      ("inflation_yoy",r"Inflación\s*\(variación %\)\s*(\w+)\s+(\d{4}).{0,180}?Interanual\s+(\d+(?:\.\d+)?)%"),
      ("core_inflation_yoy",r"Inflación subyacente\s*(\w+)\s+(\d{4}).{0,180}?Interanual\s+(\d+(?:\.\d+)?)%"),
      ("policy_rate",r"Tasa de política monetaria\s*(\w+)\s+(\d{4})\s+(\d+(?:\.\d+)?)%"),
      ("imae_yoy",r"IMAE original.*?(\w+)\s+(\d{4})\s+(\d+(?:\.\d+)?)%")
    ]
    for k,pat in patterns:
        m=re.search(pat,text,re.I|re.S)
        if m and m.group(1).lower() in MONTHS and setv(p,k,float(m.group(3)),per(m.group(1),m.group(2))):u.append(k)

    m=re.search(r"Tasas de interés\s*\(promedio ponderado\)\s*(\w+)\s+(\d{4}).{0,300}?Interbancaria\s+(\d+(?:\.\d+)?)%.{0,140}?Activa\s*B\.?M\.?\s+(\d+(?:\.\d+)?)%.{0,140}?Pasiva\s*B\.?M\.?\s+(\d+(?:\.\d+)?)%",text,re.I|re.S)
    if m and m.group(1).lower() in MONTHS:
        period=per(m.group(1),m.group(2))
        for k,v in [("interbank_rate",m.group(3)),("active_rate",m.group(4)),("passive_rate",m.group(5))]:
            if setv(p,k,float(v),period):u.append(k)

    m=re.search(r"Préstamos privados.*?(\w+)\s+(\d{4}).{0,200}?Mon\.?\s*nacional\s+(\d+(?:\.\d+)?)%.{0,120}?Total\s+(\d+(?:\.\d+)?)%",text,re.I|re.S)
    if m and m.group(1).lower() in MONTHS and setv(p,"private_credit_yoy",float(m.group(4)),per(m.group(1),m.group(2))):u.append("private_credit_yoy")

    m=re.search(r"Reservas internacionales\s*\(en US\$ MM\)\s*(\w+)\s+(\d{4}).{0,180}?Brutas\s+\$?([\d,]+(?:\.\d+)?).{0,100}?Netas\s+\$?([\d,]+(?:\.\d+)?)",text,re.I|re.S)
    if m and m.group(1).lower() in MONTHS:
        period=per(m.group(1),m.group(2))
        for k,v in [("gross_reserves",m.group(3)),("net_reserves",m.group(4))]:
            if setv(p,k,num(v),period):u.append(k)

    m=re.search(r"Producto Interno Bruto\s*\(variación % interanual\).*?Ene-Dic\s+(\d{4})\s+(\d+(?:\.\d+)?)%",text,re.I|re.S)
    if m and setv(p,"gdp_growth",float(m.group(2)),m.group(1)):u.append("gdp_growth")

    m=re.search(r"Cuenta corriente\s*\(como % del PIB\).*?2025\s+(-?\d+(?:\.\d+)?)%",text,re.I|re.S)
    if m and setv(p,"current_account_gdp",float(m.group(1)),"2025"):u.append("current_account_gdp")

    # Maintain a compact monthly history used by the dashboard charts.
    hist_path = ROOT / "data" / "history.json"
    try:
        hist = json.loads(hist_path.read_text(encoding="utf-8")) if hist_path.exists() else {}
    except Exception:
        hist = {}

    dgii_updated = update_dgii(p, hist)
    hacienda_updated = update_hacienda(p, hist)
    update_fiscal_history(p, hist)
    watch(p, hist)
    p["as_of"]=datetime.now(timezone.utc).date().isoformat()
    p["status"]="live-bcrd-dgii-hacienda-v2"
    p["automation_check_utc"]=datetime.now(timezone.utc).isoformat(timespec="seconds")
    p["updated_fields"]=u
    p.setdefault("source_health",{})["bcrd"]={"ok":True,"url":BCRD,"http_status":r.status_code,"updated_fields":u}
    hist.setdefault("monthly", {})
    for key in ["active_rate","passive_rate","interbank_rate","policy_rate","inflation_yoy",
                "core_inflation_yoy","imae_yoy","gross_reserves","net_reserves","private_credit_yoy"]:
        x = p.get("series", {}).get(key)
        if x and x.get("period") and x.get("value") is not None:
            hist["monthly"].setdefault(key, {})
            hist["monthly"][key][x["period"]] = x["value"]
    hist.setdefault("daily", {})
    for key in ["usd_dop_buy","usd_dop_sell"]:
        x = p.get("series", {}).get(key)
        if x and x.get("period") and x.get("value") is not None:
            hist["daily"].setdefault(key, {})
            hist["daily"][key][x["period"]] = x["value"]
    hist["updated_utc"] = p["automation_check_utc"]
    hist_path.write_text(json.dumps(hist, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    DATA.write_text(json.dumps(p,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    all_updated = u + dgii_updated + hacienda_updated
    print("Updated:", ", ".join(all_updated) if all_updated else "none")

if __name__=="__main__": main()
